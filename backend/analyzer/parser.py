import re
from pathlib import Path
from pypdf import PdfReader
from docx import Document


SECTION_ALIASES = {
    "summary": [
        "summary",
        "professional summary",
        "profile",
        "career summary",
        "objective",
        "professional profile",
    ],
    "skills": [
        "skills",
        "technical skills",
        "core skills",
        "technical expertise",
        "technologies",
        "technical competencies",
    ],
    "experience": [
        "experience",
        "work experience",
        "professional experience",
        "employment history",
        "career history",
    ],
    "education": [
        "education",
        "academic background",
        "academic qualifications",
        "qualifications",
    ],
    "projects": [
        "projects",
        "personal projects",
        "academic projects",
        "key projects",
    ],
    "certifications": [
        "certifications",
        "certificates",
        "professional certifications",
    ],
    "achievements": [
        "achievements",
        "accomplishments",
        "awards",
    ],
}


def clean_extraction_artifacts(text):
    """
    Repair common PDF text-extraction artifacts.

    Examples:
        "T ata" -> "Tata"
        "c onfiguration" -> "configuration"
        "servicedelays" -> "service delays"
        "compl eteness" -> "completeness"
    """

    if not text:
        return ""

    # Remove spaces accidentally inserted inside common words.
    replacements = {
        r"\bT\s+ata\b": "Tata",
        r"\bc\s+onfiguration\b": "configuration",
        r"\bservicedelays\b": "service delays",
        r"\bdocument\s+ation\b": "documentation",
        r"\bcompl\s+eteness\b": "completeness",
        r"\bprov\s+iding\b": "providing",
        r"\bprov\s+ide\b": "provide",
        r"\bproduc\s+tion\b": "production",
        r"\bconfigur\s+ation\b": "configuration",
        r"\bautomat\s+e\b": "automate",
        r"\banomalydetection\b": "anomaly detection",
        r"\bevery Git push\b": "every Git push",
        r"\bonevery\b": "on every",
        r"\bS\s+OC2\b": "SOC2",
        r"\bSS AE18\b": "SSAE18",
        r"\bA I\b": "AI",
        r"\bM FA\b": "MFA",
        r"\bR CA\b": "RCA",
        r"\bI a c\b": "IaC",
        r"\bIa c\b": "IaC",
    }

    for pattern, replacement in replacements.items():
        text = re.sub(
            pattern,
            replacement,
            text,
            flags=re.IGNORECASE,
        )

    # Remove spaces immediately before punctuation.
    text = re.sub(
        r"\s+([,.;:])",
        r"\1",
        text,
    )

    # Normalize hyphen spacing.
    text = re.sub(
        r"\s*-\s*",
        "-",
        text,
    )

    return text


def clean_text(text):
    if not text:
        return ""

    text = text.replace("\x00", " ")
    text = text.replace("\r", "\n")

    text = clean_extraction_artifacts(
        text
    )

    lines = []

    for line in text.splitlines():
        line = re.sub(
            r"[ \t]+",
            " ",
            line,
        ).strip()

        if line:
            lines.append(line)

    return "\n".join(lines)


def extract_pdf_text(file_bytes):
    try:
        reader = PdfReader(file_bytes)
    except Exception as exc:
        raise ValueError(
            f"Unable to read PDF file: {exc}"
        )

    pages = []

    for page in reader.pages:
        page_text = page.extract_text() or ""
        pages.append(page_text)

    return clean_text(
        "\n".join(pages)
    )


def extract_docx_text(file_bytes):
    try:
        document = Document(file_bytes)
    except Exception as exc:
        raise ValueError(
            f"Unable to read DOCX file: {exc}"
        )

    lines = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()

        if text:
            lines.append(text)

    for table in document.tables:
        for row in table.rows:
            values = []

            for cell in row.cells:
                value = cell.text.strip()

                if value:
                    values.append(value)

            if values:
                lines.append(
                    " | ".join(values)
                )

    return clean_text(
        "\n".join(lines)
    )


def extract_text(file_bytes, filename):
    extension = Path(
        filename
    ).suffix.lower()

    if extension == ".pdf":
        return extract_pdf_text(
            file_bytes
        )

    if extension == ".docx":
        return extract_docx_text(
            file_bytes
        )

    if extension == ".txt":
        return clean_text(
            file_bytes.decode(
                "utf-8",
                errors="ignore",
            )
        )

    raise ValueError(
        "Unsupported file type. "
        "Please upload PDF, DOCX, or TXT."
    )


def normalize_heading(text):
    value = re.sub(
        r"[^a-zA-Z ]",
        "",
        text.lower(),
    ).strip()

    value = re.sub(
        r"\s+",
        " ",
        value,
    )

    return value


def detect_section_heading(line):
    normalized = normalize_heading(
        line
    )

    for section, aliases in (
        SECTION_ALIASES.items()
    ):
        if normalized in aliases:
            return section

    return None


def extract_sections(text):
    sections = {}
    current_section = "header"

    sections[current_section] = []

    for line in text.splitlines():
        heading = detect_section_heading(
            line
        )

        if heading:
            current_section = heading

            if (
                current_section
                not in sections
            ):
                sections[current_section] = []

            continue

        sections.setdefault(
            current_section,
            [],
        ).append(line)

    cleaned = {}

    for name, lines in sections.items():
        value = "\n".join(
            line.strip()
            for line in lines
            if line.strip()
        )

        if value:
            cleaned[name] = value

    return cleaned


def get_lines(text):
    return [
        line.strip()
        for line in text.splitlines()
        if line.strip()
    ]